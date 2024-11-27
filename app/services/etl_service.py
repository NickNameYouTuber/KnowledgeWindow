from app.config.database import VectorizedKnowledgeBase
from app.connectors.csv_connector import read_csv_file
from app.connectors.txt_connector import read_txt_file
from app.repositories.knowledge_base_repository import create_knowledge_base
from sqlalchemy.orm import Session
from app.schemas.knowledge_base import KnowledgeBaseCreate
from docx import Document
import openpyxl
import PyPDF2

from app.services.vectorize_service import text_to_vector


def process_txt_file(file_path: str, db: Session):
    content = read_txt_file(file_path)
    vector = text_to_vector(content)  # Преобразуем контент в вектор
    knowledge_base_entry = VectorizedKnowledgeBase(
        title=file_path.split("/")[-1],
        content=content,
        vector=vector
    )
    db.add(knowledge_base_entry)
    db.commit()

def process_csv_file(file_path: str, db: Session):
    df = read_csv_file(file_path)
    content = df.to_string(index=False)
    vector = text_to_vector(content)  # Преобразуем контент в вектор
    knowledge_base_entry = VectorizedKnowledgeBase(
        title=file_path.split("/")[-1],
        content=content,
        vector=vector
    )
    db.add(knowledge_base_entry)
    db.commit()

def process_docx_file(file_path, db):
    doc = Document(file_path)
    content = "\n".join([para.text for para in doc.paragraphs])
    vector = text_to_vector(content)  # Преобразуем контент в вектор
    knowledge_base_entry = VectorizedKnowledgeBase(
        title=file_path.split("/")[-1],
        content=content,
        vector=vector
    )
    db.add(knowledge_base_entry)
    db.commit()

def process_xlsx_file(file_path, db):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    content = "\n".join([f"{cell.value}" for row in sheet.iter_rows() for cell in row])
    vector = text_to_vector(content)  # Преобразуем контент в вектор
    knowledge_base_entry = VectorizedKnowledgeBase(
        title=file_path.split("/")[-1],
        content=content,
        vector=vector
    )
    db.add(knowledge_base_entry)
    db.commit()

def process_pdf_file(file_path, db):
    pdf_reader = PyPDF2.PdfFileReader(file_path)
    content = "\n".join([pdf_reader.getPage(i).extract_text() for i in range(pdf_reader.numPages)])
    vector = text_to_vector(content)  # Преобразуем контент в вектор
    knowledge_base_entry = VectorizedKnowledgeBase(
        title=file_path.split("/")[-1],
        content=content,
        vector=vector
    )
    db.add(knowledge_base_entry)
    db.commit()

def process_md_file(file_path, db):
    with open(file_path, "r") as f:
        content = f.read()
    vector = text_to_vector(content)  # Преобразуем контент в вектор
    knowledge_base_entry = VectorizedKnowledgeBase(
        title=file_path.split("/")[-1],
        content=content,
        vector=vector
    )
    db.add(knowledge_base_entry)
    db.commit()